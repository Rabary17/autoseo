import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

wb = openpyxl.Workbook()

ws0 = wb.active
ws0.title = 'Légende'
ws0['A1'] = "Fichier de suivi — Mots-clés & pages — Niche Auto & mobilité"
ws0['A1'].font = Font(name='Arial', size=14, bold=True)
ws0.merge_cells('A1:D1')

ws0['A3'] = "Colonnes de l'onglet Suivi"
ws0['A3'].font = Font(name='Arial', bold=True)
cols_doc = [
    ('mot_cle_principal', "Le mot-clé de tête du cluster (celui qui définit l'URL cible)."),
    ('variantes', 'Les autres mots-clés du même cluster, séparés par ";".'),
    ('silo', 'Silo du cocon (ex: Entretien & révision).'),
    ('sous_cocon', 'Sous-cocon (ex: Vidange & filtres).'),
    ('intention', 'Info / Commercial / Transactionnel.'),
    ('volume_estime', 'Volume de recherche réel (source Haloscan), jamais estimé à la main.'),
    ('concurrence', "Concurrence Haloscan pondérée par volume (0-1, plus bas = moins de concurrence), 0.5 si inconnue."),
    ('score_opportunite', "volume_estime * (1 - concurrence) — sert au tri de priorité de production, pas le volume seul (2026-08-26)."),
    ('url_cible', 'URL prévue ou publiée qui couvre ce cluster.'),
    ('auteur', "Persona auteur WordPress assigné (A à F, voir skills/wordpress-publication.md section 4)."),
    ('statut', 'à faire / en rédaction / programmé / publié / à réécrire. "programmé" = post_status WordPress future (date fixée, pas encore en ligne).'),
    ('date_publication', "Date réelle ou programmée de mise en ligne (AAAA-MM-JJ) — correspond au post_date WordPress."),
]
r = 4
for name, desc in cols_doc:
    ws0.cell(row=r, column=1, value=name).font = Font(name='Arial', bold=True)
    ws0.cell(row=r, column=2, value=desc).font = Font(name='Arial')
    r += 1

ws0['A' + str(r + 1)] = "Règle anti-cannibalisation : un mot-clé principal = une seule ligne = une seule URL. Vérifier avant tout ajout."
ws0['A' + str(r + 1)].font = Font(name='Arial', italic=True, color='C0392B')
ws0.column_dimensions['A'].width = 22
ws0.column_dimensions['B'].width = 90

ws = wb.create_sheet('Suivi')
headers = ['mot_cle_principal', 'variantes', 'silo', 'sous_cocon', 'intention', 'volume_estime',
           'concurrence', 'score_opportunite', 'url_cible', 'auteur', 'statut', 'date_publication']
header_fill = PatternFill(start_color='1E2535', end_color='1E2535', fill_type='solid')
header_font = Font(name='Arial', bold=True, color='FFFFFF')
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=i, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center')

example = [
    'vidange prix moyen',
    'vidange prix voiture;vidange tarif;vidange combien coute',
    'Entretien & révision',
    'Vidange & filtres',
    'Info',
    94,
    0.22,
    73.32,
    '/entretien-revision/vidange-filtres/vidange-prix-moyen/',
    'A — Mécanique & technique',
    'à faire',
    '',
]
example_fill = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')
for i, v in enumerate(example, start=1):
    c = ws.cell(row=2, column=i, value=v)
    c.font = Font(name='Arial', italic=True)
    c.fill = example_fill
ws.cell(row=2, column=1).comment = Comment(
    "Ligne d'exemple — à remplacer/supprimer une fois les vraies données Haloscan collectées.",
    'autoseo',
)

widths = [30, 45, 22, 22, 14, 14, 12, 16, 45, 24, 14, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = 'A2'

wb.save('data/keywords/tracking-mots-cles.xlsx')
print('saved')
