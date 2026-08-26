# Peuple data/keywords/tracking-mots-cles.xlsx (onglet "Suivi") avec les vraies données Haloscan
# collectées par scripts/fetch-keywords.js, silo par silo (uniquement ceux déjà traités en P1),
# PLUS les candidats programmatiques des moteurs (scripts/expand-moteurs.py) déjà validés par
# Haloscan (scripts/validate-moteurs-haloscan.js, statut 'retenu' dans moteurs-candidats.json).
#
# Une ligne = un cluster = un seed du plan (mot_cle_principal) + ses variantes enrichies
# (match/questions/related, filtrées des marques/navigation) — jamais une ligne par variante,
# pour respecter la règle anti-cannibalisation (skills/wordpress-publication.md section 5 et
# skills/gestion-de-projet.md) : un cluster = une seule URL cible.
#
# volume_estime = volume cumulé du cluster (mot-clé principal + variantes retenues), pas le seul
# volume du mot-clé de tête, pour permettre un tri par priorité réelle (volume bon / difficulté faible).
#
# CORRIGÉ le 2026-08-26 (demande explicite de l'utilisateur, priorité qualité > quantité après
# un déclassement attribué à un profil "keyword farm") : les variantes étaient piochées dans
# `candidates` (mots-clés "liés" bruts renvoyés par Haloscan, SANS aucun score de pertinence),
# pas dans `results` (les mêmes mots-clés mais après le passage de similarité sémantique de
# Haloscan, avec un vrai score `similarity`). Constaté en conditions réelles : le cluster
# "assurance voiture sans permis" (volume réel 210) affichait un `volume_estime` de 1 317 846,
# parce que `candidates` incluait "caisse primaire d'assurance maladie" (611 818/mois) et
# "direct assurance" (488 200/mois) — deux requêtes sans aucun rapport thématique, seul le mot
# "assurance" étant commun. Sur 188 seeds audités, 30 (16%) avaient ce défaut. Cette contamination
# touchait aussi la colonne `variantes` elle-même, pas seulement le volume affiché.
#
# IMPORTANT : ce script REGÉNÈRE intégralement l'onglet "Suivi" à chaque exécution (source de
# vérité = seeds.json éditorial + moteurs-candidats.json 'retenu', jamais l'xlsx lui-même). Pour
# ne jamais perdre la progression de publication déjà faite par le pipeline autopublish (statut
# 'en rédaction'/'programmé'/'publié', url_cible, date_publication écrits par run.js), l'ancien
# état de chaque mot_cle_principal est capturé AVANT régénération et réappliqué après si le
# cluster avait progressé au-delà de 'à faire' — voir preserve_publish_state().
#
# Usage : python scripts/populate-tracking-xlsx.py [--niche <id>]
#
# 2026-07-20 (industrialisation) : AUTHOR_MAP/SILO_ORDER étaient codés en dur en Python,
# spécifiques à "Auto & mobilité" — ajouter un silo exigeait d'éditer ce fichier. Remplacés
# par la config déclarative config/niches/<id>/niche.json (scripts/lib_py/niche_config.py).
# --niche par défaut 'auto-mobilite', dont data_dir='data' pointe sur la racine actuelle :
# comportement 100% identique à avant pour qui ne passe pas ce flag.

import argparse
import json
import os
import re
import sys
import unicodedata

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lib_py'))
import niche_config

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

HEADERS = ['mot_cle_principal', 'variantes', 'silo', 'sous_cocon', 'intention',
           'volume_estime', 'concurrence', 'score_opportunite',
           'url_cible', 'auteur', 'statut', 'date_publication']

INTENT_LABELS = {'I': 'Info', 'C': 'Commercial', 'T': 'Transactionnel'}

MAX_VARIANTES = 15

# Seuil de similarité Haloscan (0-1) sous lequel un résultat de `results` n'est
# plus considéré assez proche du mot-clé de tête pour compter dans le cluster —
# généreux par construction : sur les échantillons audités le 2026-08-26, même
# le résultat le moins similaire d'un cluster propre restait > 0.75.
SIMILARITY_MIN = 0.5
# Concurrence neutre appliquée quand Haloscan n'a renvoyé aucune valeur pour un
# candidat retenu — ne pas confondre "concurrence inconnue" avec "nulle" (qui
# survaloriserait à tort) ni "maximale" (qui dévaloriserait à tort).
DEFAULT_COMPETITION = 0.5


def slug(name):
    n = unicodedata.normalize('NFD', name.lower())
    n = ''.join(c for c in n if unicodedata.category(c) != 'Mn')
    n = re.sub(r'[^a-z0-9]+', '-', n).strip('-')
    return n


def load_silo_rows(data_dir, author_map, silo_name):
    path = os.path.join(data_dir, 'keywords', f'{slug(silo_name)}.json')
    if not os.path.exists(path):
        return []
    with open(path, encoding='utf-8') as f:
        data = json.load(f)

    def as_num(v):
        return v if isinstance(v, (int, float)) else None

    rows = []
    for keyword, entry in data.get('seeds', {}).items():
        results = entry.get('results') or []
        relevant = [r for r in results
                    if isinstance(r.get('similarity'), (int, float)) and r['similarity'] >= SIMILARITY_MIN]
        top_variantes = relevant[:MAX_VARIANTES]

        seed_volume = as_num(entry.get('volume')) or 0
        variante_volumes = [as_num(c.get('volume')) or 0 for c in top_variantes]
        volume_cluster = seed_volume + sum(variante_volumes)

        # Concurrence pondérée par volume : un candidat à fort volume pèse plus
        # dans le score qu'un candidat marginal, plutôt qu'une moyenne simple.
        weighted = [(as_num(c.get('competition')), as_num(c.get('volume')) or 0)
                    for c in top_variantes if as_num(c.get('competition')) is not None]
        if weighted:
            total_weight = sum(w for _, w in weighted) or len(weighted)
            concurrence = sum((c or 0) * (w or 1) for c, w in weighted) / total_weight
        else:
            concurrence = DEFAULT_COMPETITION
        score_opportunite = round(volume_cluster * (1 - concurrence), 2)

        rows.append({
            'mot_cle_principal': keyword,
            'variantes': ';'.join(c['keyword'] for c in top_variantes),
            'silo': silo_name,
            'sous_cocon': entry.get('sub', ''),
            'intention': INTENT_LABELS.get(entry.get('intent'), ''),
            'volume_estime': volume_cluster,
            'concurrence': round(concurrence, 3),
            'score_opportunite': score_opportunite,
            'url_cible': '',
            'auteur': author_map.get(silo_name, ''),
            'statut': 'à faire',
            'date_publication': '',
        })
    # priorité : score d'opportunité décroissant (fort volume + faible concurrence
    # d'abord — 2026-08-26, remplace le tri par volume brut seul), à l'intérieur
    # d'un même silo.
    rows.sort(key=lambda r: r['score_opportunite'], reverse=True)
    return rows


def load_moteurs_rows_by_silo(moteurs_json_path, author_map):
    """Candidats programmatiques (scripts/expand-moteurs.py) déjà validés par Haloscan
    (statut 'retenu' écrit par scripts/validate-moteurs-haloscan.js), groupés par silo.
    Fichier absent ou vide -> aucun impact (rétro-compatible avec le comportement d'origine,
    100 % éditorial)."""
    if not os.path.exists(moteurs_json_path):
        return {}
    with open(moteurs_json_path, encoding='utf-8') as f:
        data = json.load(f)

    by_silo = {}
    for c in data.get('candidats', []):
        if c.get('statut') != 'retenu':
            continue
        # Pas de donnée de concurrence pour les candidats programmatiques (validés
        # par scripts/validate-moteurs-haloscan.js, qui ne la capture pas encore) —
        # concurrence neutre, comportement équivalent au tri par volume seul.
        volume_estime = c.get('volume_estime') or 0
        by_silo.setdefault(c['silo'], []).append({
            'mot_cle_principal': c['mot_cle_principal'],
            'variantes': c['variantes'],
            'silo': c['silo'],
            'sous_cocon': c['sous_cocon'],
            'intention': c['intention'],
            'volume_estime': volume_estime,
            'concurrence': DEFAULT_COMPETITION,
            'score_opportunite': round(volume_estime * (1 - DEFAULT_COMPETITION), 2),
            'url_cible': '',
            'auteur': author_map.get(c['silo'], ''),
            'statut': 'à faire',
            'date_publication': '',
        })
    return by_silo


def build_all_rows(data_dir, moteurs_json_path, author_map, silo_order):
    moteurs_by_silo = load_moteurs_rows_by_silo(moteurs_json_path, author_map)
    all_rows = []
    processed_silos = []
    moteurs_counts = {}
    for silo_name in silo_order:
        editorial_rows = load_silo_rows(data_dir, author_map, silo_name)
        moteurs_rows = moteurs_by_silo.get(silo_name, [])
        rows = editorial_rows + moteurs_rows
        # priorité : score d'opportunité décroissant sur l'ensemble éditorial + programmatique,
        # pas seulement au sein de chaque source (skills/seo.md section 3).
        rows.sort(key=lambda r: r['score_opportunite'], reverse=True)
        if rows:
            processed_silos.append((silo_name, len(editorial_rows), len(moteurs_rows)))
            all_rows.extend(rows)
        if moteurs_rows:
            moteurs_counts[silo_name] = len(moteurs_rows)
    return all_rows, processed_silos, moteurs_counts


def capture_publish_state(ws):
    """Avant régénération : capture l'état de publication déjà avancé (par le pipeline
    autopublish, run.js) pour chaque mot_cle_principal, afin de ne jamais le perdre en
    reconstruisant l'onglet à partir des sources (seeds.json + moteurs-candidats.json) —
    ce script régénère TOUT à chaque run, donc sans cette capture un run relancé après le
    début de la publication réelle écraserait silencieusement statut/url_cible/date.

    Lecture par NOM de colonne (pas par index fixe) depuis le 2026-08-26 : la mise en
    colonnes a changé (ajout concurrence/score_opportunite) et un index en dur se
    désynchronise silencieusement à chaque futur changement de colonnes."""
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col = {name: i for i, name in enumerate(header_row)}
    state = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        mot_cle = row[col['mot_cle_principal']]
        statut = row[col['statut']]
        url_cible = row[col['url_cible']]
        date_pub = row[col['date_publication']]
        if mot_cle and statut and statut != 'à faire':
            state[mot_cle] = {'statut': statut, 'url_cible': url_cible, 'date_publication': date_pub}
    return state


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--niche', default=niche_config.DEFAULT_NICHE)
    args = parser.parse_args()

    niche = niche_config.load_niche(args.niche)
    data_dir = niche_config.data_path(niche)
    xlsx_path = niche_config.data_path(niche, 'keywords', 'tracking-mots-cles.xlsx')
    moteurs_json_path = niche_config.data_path(niche, 'keywords', 'moteurs-candidats.json')
    author_map = niche_config.author_map(niche)
    silo_order = niche_config.silo_order(niche)

    all_rows, processed_silos, moteurs_counts = build_all_rows(data_dir, moteurs_json_path, author_map, silo_order)

    if not os.path.exists(xlsx_path):
        raise SystemExit(f'{xlsx_path} introuvable — lancer scripts/build-tracking-xlsx.py une première fois.')

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb['Suivi']

    old_publish_state = capture_publish_state(ws)

    # Resynchronise la ligne d'en-tête sur HEADERS (2026-08-26) : ce script ne réécrivait
    # jusqu'ici JAMAIS la ligne 1, seulement les données à partir de la ligne 2 — correct
    # tant que HEADERS ne change pas, mais silencieusement destructeur le jour où une colonne
    # est ajoutée (ici concurrence/score_opportunite) : les valeurs continuent de s'écrire
    # dans l'ordre de HEADERS, mais les libellés d'en-tête restent ceux de l'ancienne liste,
    # donc tout ce qui suit la nouvelle colonne se retrouve décalé d'une case par rapport à
    # son libellé — constaté en testant sur une copie avant d'écrire sur le fichier réel.
    header_fill = PatternFill(start_color='1E2535', end_color='1E2535', fill_type='solid')
    header_font = Font(name='Arial', bold=True, color='FFFFFF')
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center')

    # Repart d'un onglet Suivi vierge (header conservé) pour éviter les doublons entre deux runs.
    ws.delete_rows(2, ws.max_row)

    r = 2
    preserved = 0
    for row in all_rows:
        prior = old_publish_state.get(row['mot_cle_principal'])
        if prior:
            row = {**row, **prior}
            preserved += 1
        for i, h in enumerate(HEADERS, start=1):
            ws.cell(row=r, column=i, value=row[h])
        r += 1

    ws.freeze_panes = 'A2'
    widths = [30, 60, 22, 22, 14, 14, 12, 16, 45, 24, 14, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Note de suivi dans l'onglet Légende : silos déjà couverts par cette génération, avec
    # le détail éditorial vs programmatique (moteurs) par silo.
    #
    # Le bloc de note précédent (marqueur ci-dessous) est SUPPRIMÉ avant d'en écrire un
    # nouveau — sans ça, chaque exécution du script empile un nouveau bloc à la suite du
    # précédent sans jamais rien effacer. Constaté en conditions réelles : 246 lignes déjà
    # accumulées dans l'onglet Légende avant ce correctif, faute d'avoir jamais été nettoyées
    # au fil des dizaines d'exécutions passées (voir STATE.md).
    legend = wb['Légende']
    MARKER = 'Dernière génération automatique (populate-tracking-xlsx.py)'
    existing_marker_row = next(
        (row for row in range(1, legend.max_row + 1) if legend.cell(row=row, column=1).value == MARKER),
        None,
    )
    if existing_marker_row is not None:
        legend.delete_rows(existing_marker_row, legend.max_row - existing_marker_row + 1)
        note_row = existing_marker_row
    else:
        note_row = legend.max_row + 2
    legend.cell(row=note_row, column=1, value=MARKER).font = Font(name='Arial', bold=True)
    r2 = note_row + 1
    for silo_name, n_edit, n_moteurs in processed_silos:
        detail = f'{n_edit} éditoriaux'
        if n_moteurs:
            detail += f' + {n_moteurs} programmatiques'
        legend.cell(row=r2, column=1, value=silo_name)
        legend.cell(row=r2, column=2, value=f'{n_edit + n_moteurs} clusters ({detail})')
        r2 += 1
    if not processed_silos:
        legend.cell(row=r2, column=1, value='(aucun silo P1 traité pour le moment)')

    wb.save(xlsx_path)
    total = len(all_rows)
    total_moteurs = sum(moteurs_counts.values())
    print(f'{total} clusters écrits dans {xlsx_path} (dont {total_moteurs} programmatiques) — '
          f'silos couverts : {", ".join(s for s, _, _ in processed_silos) or "aucun"}')
    if preserved:
        print(f'{preserved} lignes ont conservé leur progression de publication déjà en cours '
              f'(statut/url_cible/date_publication préservés).')


if __name__ == '__main__':
    main()
